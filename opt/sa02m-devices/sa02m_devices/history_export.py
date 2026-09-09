"""Export tables → TSV / xlsx (openpyxl when present, minimal OOXML otherwise)."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from sa02m_devices.history_metrics import HISTORY_GROUPS, METRICS
from sa02m_devices.history_query import history, history_carel_batch
from sa02m_devices.history_ranges import (
    _TZ,
    _normalize_range,
    _now_local,
    _range_slug,
    export_bucket_s,
    range_label_ru,
    resolve_time_range,
)


def _fmt_export_ts(ts_s: float, bucket_s: float) -> str:
    dt = datetime.fromtimestamp(float(ts_s), tz=_TZ)
    if bucket_s >= 86400:
        return dt.strftime("%Y-%m-%d")
    if bucket_s >= 3600:
        return dt.strftime("%Y-%m-%d %H:00")
    return dt.strftime("%Y-%m-%d %H:%M")


def _export_bucket_label(bucket_s: float) -> str:
    if bucket_s <= 60:
        return "1 мин"
    if bucket_s <= 1800:
        return "30 мин"
    if bucket_s <= 3600:
        return "1 ч"
    return "1 сут"


# Title of a multi-metric export by group; a single metric is titled by its label.
_GROUP_TITLES: dict[str, str] = {"climate": "Климат", "energy": "Энергия", "ahu": "Carel AHU"}


def _export_col_title(metric_id: str, field: str) -> str:
    meta = METRICS[metric_id]
    lab = str(meta["labels"].get(field) or field)
    unit = str(meta.get("unit") or "").strip()
    return f"{lab}, {unit}" if unit else lab


def collect_export_table(
    range_key: str = "1h",
    *,
    metric_id: str | None = None,
    group: str | None = None,
    device_id: str | None = None,
    path: Path | None = None,
) -> dict[str, Any]:
    """Собрать одну таблицу (время × колонки) для экспорта."""
    range_key = _normalize_range(range_key)
    bucket = export_bucket_s(range_key)
    t0, t1, _ = resolve_time_range(range_key)
    did = (device_id or "").strip()

    if group and group in HISTORY_GROUPS:
        metric_ids = list(HISTORY_GROUPS[group])
    elif metric_id and metric_id in METRICS:
        metric_ids = [metric_id]
    else:
        return {
            "ok": False,
            "error": "укажите metric=… или group=climate|energy|ahu",
            "headers": ["Время"],
            "rows": [],
            "device_id": did,
            "range": range_key,
            "bucket_s": bucket,
            "t0": t0,
            "t1": t1,
            "metric_ids": [],
            "kind": "device",
        }

    col_fields: list[tuple[str, str]] = []  # (metric_id, field)
    for mid in metric_ids:
        for field in METRICS[mid]["fields"]:
            col_fields.append((mid, field))

    by_ts: dict[int, dict[str, float]] = {}
    for mid in metric_ids:
        one = history(
            mid,
            range_key,
            path=path,
            device_id=did or None,
            bucket_s=bucket,
            # The counter exports its bucket max (a monotonic total, averaged,
            # under-reads); every other metric takes its METRICS default.
            agg="max" if mid == "energy_kwh_import" else None,
        )
        if not did and one.get("device_id"):
            did = str(one["device_id"])
        for ser in one.get("series") or []:
            field = str(ser.get("field") or "")
            if not field:
                continue
            key = f"{mid}:{field}"
            for ts_ms, val in ser.get("points") or []:
                try:
                    by_ts.setdefault(int(ts_ms), {})[key] = float(val)
                except (TypeError, ValueError):
                    continue

    headers = ["Время"] + [_export_col_title(m, f) for m, f in col_fields]
    rows: list[list[Any]] = []
    for ts_ms in sorted(by_ts):
        cells: list[Any] = [_fmt_export_ts(ts_ms / 1000.0, bucket)]
        vals = by_ts[ts_ms]
        for mid, field in col_fields:
            v = vals.get(f"{mid}:{field}")
            cells.append(None if v is None else float(v))
        rows.append(cells)

    # kind = the metric set's device (dtv / ce / carel), `mixed` across devices.
    devices = {str(METRICS[m]["device"]) for m in metric_ids}
    kind = devices.pop() if len(devices) == 1 else "mixed"

    return {
        "ok": True,
        "error": "",
        "headers": headers,
        "rows": rows,
        "device_id": did,
        "range": range_key,
        "bucket_s": bucket,
        "t0": t0,
        "t1": t1,
        "metric_ids": metric_ids,
        "kind": kind,
        "title": (
            METRICS[metric_ids[0]]["label"]
            if len(metric_ids) == 1
            else _GROUP_TITLES.get(str(group), "Данные")
        ),
    }


def collect_export_table_carel(
    range_key: str = "1h",
    *,
    device_id: str | None = None,
    path: Path | None = None,
) -> dict[str, Any]:
    """Таблица экспорта Carel (время × метрики) — форма как collect_export_table_mr().

    Колонки берутся из ПРИСУТСТВУЮЩИХ метрик (адаптер уже отбросил неснятые
    пробы), а не из всей группы `ahu`: неукомплектованный датчик не должен
    добавлять в Excel пустой столбец — то же правило, что и на карточке (§7).
    """
    range_key = _normalize_range(range_key)
    bucket = export_bucket_s(range_key)
    batch = history_carel_batch(device_id, range_key, path=path, bucket_s=bucket)
    did = str(batch.get("device_id") or (device_id or "").strip())
    t0 = float(batch.get("t0") or 0.0)
    t1 = float(batch.get("t1") or 0.0)

    col_fields: list[str] = []
    col_titles: list[str] = []
    by_ts: dict[int, dict[str, float]] = {}
    for metric in batch.get("metrics") or []:
        unit = str(metric.get("unit") or "").strip()
        for ser in metric.get("series") or []:
            field = str(ser.get("field") or "")
            if not field:
                continue
            label = str(ser.get("label") or field)
            col_fields.append(field)
            col_titles.append(f"{label}, {unit}" if unit else label)
            for ts_ms, val in ser.get("points") or []:
                try:
                    by_ts.setdefault(int(ts_ms), {})[field] = float(val)
                except (TypeError, ValueError):
                    continue

    headers = ["Время"] + col_titles
    rows: list[list[Any]] = []
    for ts_ms in sorted(by_ts):
        cells: list[Any] = [_fmt_export_ts(ts_ms / 1000.0, bucket)]
        vals = by_ts[ts_ms]
        for field in col_fields:
            v = vals.get(field)
            cells.append(None if v is None else float(v))
        rows.append(cells)

    return {
        "ok": True,
        "error": "",
        "headers": headers,
        "rows": rows,
        "device_id": did,
        "range": range_key,
        "bucket_s": bucket,
        "t0": t0,
        "t1": t1,
        "metric_ids": col_fields,
        "kind": "carel",
        "title": _GROUP_TITLES["ahu"],
    }


def _xml_escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _export_xlsx_minimal(table: dict[str, Any], filename: str) -> bytes:
    """OOXML .xlsx без openpyxl (достаточно для Excel/LibreOffice)."""
    import zipfile
    from io import BytesIO

    def cell_ref(row: int, col: int) -> str:
        n = col
        letters = ""
        while n:
            n, rem = divmod(n - 1, 26)
            letters = chr(65 + rem) + letters
        return f"{letters}{row}"

    def inline_cell(row: int, col: int, val: Any) -> str:
        ref = cell_ref(row, col)
        if val is None:
            return f'<c r="{ref}"/>'
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            return f'<c r="{ref}"><v>{val}</v></c>'
        text = _xml_escape(str(val))
        return (
            f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'
        )

    meta_rows = [
        ("Устройство", table.get("device_id") or "—"),
        ("Метрика", table.get("title") or "—"),
        ("Строк", str(len(table.get("rows") or []))),
    ]
    headers = list(table.get("headers") or ["Время"])
    data_rows = list(table.get("rows") or [])
    sheet_rows: list[str] = []
    r = 1
    for k, v in meta_rows:
        sheet_rows.append(
            f'<row r="{r}">'
            f"{inline_cell(r, 1, k)}{inline_cell(r, 2, v)}"
            f"</row>"
        )
        r += 1
    r = 5
    sheet_rows.append(
        f'<row r="{r}">'
        + "".join(inline_cell(r, c, h) for c, h in enumerate(headers, start=1))
        + "</row>"
    )
    if not data_rows:
        r = 6
        sheet_rows.append(
            f'<row r="{r}">{inline_cell(r, 1, "(нет точек за период)")}</row>'
        )
    else:
        for row in data_rows:
            r += 1
            sheet_rows.append(
                f'<row r="{r}">'
                + "".join(
                    inline_cell(r, c, val) for c, val in enumerate(row, start=1)
                )
                + "</row>"
            )

    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"<sheetData>{''.join(sheet_rows)}</sheetData></worksheet>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    wb_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Данные" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("xl/workbook.xml", wb_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    _ = filename  # kept for call-site symmetry
    return buf.getvalue()


def export_xlsx(
    range_key: str = "1h",
    *,
    metric_id: str | None = None,
    group: str | None = None,
    device_id: str | None = None,
    kind: str | None = None,
    path: Path | None = None,
) -> tuple[bytes, str]:
    """Выгрузка в Excel (.xlsx) с таблицей. Возвращает (bytes, filename)."""
    from sa02m_devices.history_mr import collect_export_table_mr
    from io import BytesIO

    if kind == "mr":
        table = collect_export_table_mr(range_key, device_id=device_id, path=path)
    elif kind == "carel":
        table = collect_export_table_carel(range_key, device_id=device_id, path=path)
    else:
        table = collect_export_table(
            range_key, metric_id=metric_id, group=group, device_id=device_id, path=path
        )
    stamp = _now_local().strftime("%Y%m%d_%H%M%S")
    mid_part = metric_id or group or (
        "ai" if kind == "mr" else "ahu" if kind == "carel" else "data"
    )
    safe_id = (str(table.get("device_id") or "device")).replace("/", "-")
    kind = str(table.get("kind") or "device")
    filename = f"{kind}_export_{safe_id}_{mid_part}_{_range_slug(range_key)}_{stamp}.xlsx"

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        return _export_xlsx_minimal(table, filename), filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"

    header_font = Font(bold=True)
    meta_font = Font(bold=True, color="334455")
    meta_fill = PatternFill("solid", fgColor="E8EEF5")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    meta_rows = [
        ("Устройство", table.get("device_id") or "—"),
        ("Метрика", table.get("title") or mid_part),
        (
            "Период",
            f"{range_label_ru(range_key)} "
            f"({_fmt_export_ts(float(table['t0']), 60)} — "
            f"{_fmt_export_ts(float(table['t1']), 60)})",
        ),
        ("Шаг", _export_bucket_label(float(table["bucket_s"]))),
        ("Строк", len(table.get("rows") or [])),
    ]
    for i, (k, v) in enumerate(meta_rows, start=1):
        ws.cell(i, 1, k).font = meta_font
        ws.cell(i, 1).fill = meta_fill
        ws.cell(i, 2, v)

    header_row = 7
    headers = list(table.get("headers") or ["Время"])
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, h)
        cell.font = header_font
        cell.alignment = center

    data_rows = list(table.get("rows") or [])
    if not data_rows:
        ws.cell(header_row + 1, 1, "(нет точек за период)")
    else:
        for r_i, row in enumerate(data_rows, start=header_row + 1):
            for c_i, val in enumerate(row, start=1):
                cell = ws.cell(r_i, c_i, val)
                cell.alignment = center
                if c_i > 1 and isinstance(val, float):
                    cell.number_format = "0.###"

    last_data_row = header_row + max(1, len(data_rows))
    last_col = max(1, len(headers))
    # Excel Table (полноценная таблица)
    if data_rows:
        ref = f"A{header_row}:{get_column_letter(last_col)}{last_data_row}"
        tab = Table(displayName="ExportData", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)

    ws.column_dimensions["A"].width = 20
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = f"A{header_row + 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), filename


def export_text(
    range_key: str = "1h",
    *,
    metric_id: str | None = None,
    group: str | None = None,
    device_id: str | None = None,
    kind: str | None = None,
    path: Path | None = None,
) -> tuple[str, str]:
    """Текстовая выгрузка (TSV) — совместимость; основной формат: export_xlsx."""
    from sa02m_devices.history_mr import collect_export_table_mr
    if kind == "mr":
        table = collect_export_table_mr(range_key, device_id=device_id, path=path)
    elif kind == "carel":
        table = collect_export_table_carel(range_key, device_id=device_id, path=path)
    else:
        table = collect_export_table(
            range_key, metric_id=metric_id, group=group, device_id=device_id, path=path
        )
    if not table.get("ok"):
        return f"# error: {table.get('error')}\n", "export_error.txt"
    headers = table["headers"]
    lines = [
        f"# устройство: {table.get('device_id') or '—'}",
        f"# метрика: {table.get('title')}",
        f"# период: {range_label_ru(range_key)}"
        f" ({_fmt_export_ts(float(table['t0']), 60)} — "
        f"{_fmt_export_ts(float(table['t1']), 60)})",
        f"# шаг: {_export_bucket_label(float(table['bucket_s']))}",
        "\t".join(str(h) for h in headers),
    ]
    for row in table.get("rows") or []:
        cells = []
        for v in row:
            if v is None:
                cells.append("")
            elif isinstance(v, float):
                cells.append(f"{v:.6g}")
            else:
                cells.append(str(v))
        lines.append("\t".join(cells))
    if not table.get("rows"):
        lines.append("# (нет точек)")
    body = "\n".join(lines) + "\n"
    stamp = _now_local().strftime("%Y%m%d_%H%M%S")
    mid_part = metric_id or group or (
        "ai" if kind == "mr" else "ahu" if kind == "carel" else "data"
    )
    safe_id = (str(table.get("device_id") or "device")).replace("/", "-")
    kind_part = str(table.get("kind") or "device")
    filename = f"{kind_part}_export_{safe_id}_{mid_part}_{_range_slug(range_key)}_{stamp}.txt"
    return body, filename
